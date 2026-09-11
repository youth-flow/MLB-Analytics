from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


ROOT = Path(r"C:\Users\yyx\Desktop\大三暑期短学期")
WEEK_DIR = ROOT / "提交材料" / "01_每周实习记录"
WORK_DIR = ROOT / ".submission_revision_work"
REFERENCE = WEEK_DIR / "第1周实习记录.docx"
NOTE = "说明：以下内容按项目实际推进顺序整理，具体日期以会议材料和实验记录为准。"


WEEK3 = [
    ("note", NOTE),
    ("heading", "一、从“球威下降”转向“控球问题”"),
    (
        "body",
        "本周开始对前两周冻结的数据做正式分析。我把今井在NPB 2025和MLB 2026的三振、保送、挥空与球速放在一起核对。两阶段K%均为27.8%，MLB Whiff%为32.3%，四缝线均速也接近NPB时期，至少不支持“球威整体消失”。更明显的变化是BB%由7.0%升至14.6%，Zone%为43.6%，首球好球率为51.7%。",
    ),
    (
        "body",
        "这使我把重点从ERA转向过程指标。只看5.29的ERA，很容易归因于联盟强度或球速；加入逐球信息后，当前更需要解决的是进区和抢0—1球数的稳定性。Zone%又不能代表位置质量，所以我继续检查四缝线落点、滑球能否形成有竞争力的坏球和不同球数下的选择，没有把“多投好球”理解成往好球区中央投。",
    ),
    ("heading", "二、球种结构与打者侧别"),
    (
        "body",
        "今井在MLB的四缝线和滑球合计占88.9%。我起初想减少滑球，但它的Whiff为40.6%、官方Run Value为+4，机械减量并不合理。回看NPB数据，我发现2025年变速球使用率为12.7%、Whiff为41.6%、xPV/100为+1.03，而MLB记录中的使用率只有2.3%。",
    ),
    (
        "body",
        "我因此把“开发第三球种”改为“先核对并试着恢复已有的变速球”。MLB样本只有29球，不能凭命中结果判断质量；其均速由84.7 mph升至86.8 mph，与四缝线的速度差由10.14 mph收窄至8.01 mph，值得检查。两个联盟的球种标签也未必一一对应，MLB还记录了伸卡球和指叉球，所以12.7%到2.3%不能全解释成真实弃投。后续需结合球速、移动和释放特征核对球形，再决定试验用量。",
    ),
    (
        "body",
        "左右打拆分给出风险信号：四缝线对左打311球，Whiff为15.4%、RV/100为-1.27；对右打229球，分别为27.0%和+1.47。但两组的球数、落点、对手和长打结果不同，不能直接写成“左打导致失效”。我只把它作为优先检查对左打位置和速度配套的依据，并保留对右打仍有效的四缝线。",
    ),
    ("heading", "三、牛棚结果的反证检查"),
    (
        "body",
        "两次牛棚登板的Zone%、首球好球率和Whiff%分别为46.1%、73.9%和40.0%，方向上有所改善。但样本仅89球、23个打席，日期、伤病恢复和对手也与角色调整同时变化，四缝线均速并未明显上升。因此，我只把牛棚看作继续观察和校准的环境，不能说角色变化使控球提高，更不能提前决定长期改打后援。",
    ),
    (
        "body",
        "我还检查了第三轮的161球。四缝线均速为95.65 mph，投手视角Run Value合计仍为正，不支持“第三轮一定因球速下降而崩盘”；但样本也不足以排除疲劳。下一周要把角色判断改成可升级、暂停和退回的负荷方案，而不是直接二选一。",
    ),
    ("heading", "四、港科大暑研：HPC4小规模试跑"),
    (
        "body",
        "港科暑研本周开始在HPC4上试跑Synthetic-oracle实验。我先用小规模数据接通生成、偏好配对、reward训练、局部策略更新和重新评价。实验以Qwen2.5-1.5B-Instruct为参考策略，冻结表示层，只训练线性reward head；策略端只更新最后一层q_proj和v_proj的rank-4 LoRA。每个prompt生成6个回答并构造15条无序回答对，供后续估计Fisher和reward moment。",
    ),
    (
        "body",
        "第一次试跑能结束，但“程序不报错”不等于索引正确。我抽取若干prompt，手工核对6个回答、15条回答对和标签，再检查数据切分，以及策略评价是否在held-out prompts上重新生成。确认reward写回和单步策略更新按顺序接通后，我才准备下一周的三随机种子正式实验。本周只完成管线试跑和排错，没有把小规模输出当作结论。",
    ),
    ("heading", "五、阶段反思与下一步"),
    (
        "body",
        "本周最大的变化，是让结论强度与证据相匹配。变速球只有29球、牛棚只有23个打席，左右打拆分也受球数和对手影响；把限制写在判断旁边，能避免把暂时信号写成稳定规律。HPC4试跑同样说明，作业正常结束只代表流程能运行，正式实验前还必须检查索引、数据隔离和评价对象。",
    ),
    ("bullet", "完成控球、球种、左右打和角色四轮分析，并形成核心诊断图和角色拆分图。"),
    ("bullet", "将“减少滑球、开发任意第三球种”修正为“保留有效滑球，核对标签后优先试验变速球”。"),
    ("bullet", "完成HPC4小规模端到端试跑和关键索引检查，未把试跑结果写成正式结论。"),
    ("bullet", "下一周确定重返先发的观察与负荷门槛，并完成三随机种子实验、结果汇总和答辩。"),
]


WEEK4 = [
    ("note", NOTE),
    ("heading", "一、阶段主要任务"),
    (
        "body",
        "最后一周，我没有继续增加新的分析维度，而是把前三周的结果收拢成一套能够执行和复查的方案。本周依次重核关键数字与反例，确定训练顺序、观察门槛和角色分支，再整理两页汇总、底稿和仓库。最后确认：今井的K%仍为27.8%、Whiff%为32.3%，四缝线均速约94.9 mph；更明显的问题是BB%由7.0%升至14.6%，Zone%为43.6%，首球好球率为51.7%。",
    ),
    (
        "body",
        "球种方面，四缝线和滑球合计占88.9%，但滑球仍有正向价值，不能机械减量。NPB 2025变速球使用率为12.7%，到MLB标签中降至2.3%，速度差也由约10.1 mph缩小到8.0 mph。结合第三周对标签口径和小样本的检查，我把它写成“恢复并验证历史球形”，而不是认定只要增加变速球就能解决问题。",
    ),
    ("heading", "二、把判断改成执行方案"),
    (
        "body",
        "我起初想直接回答“继续先发还是转入牛棚”，但两次牛棚登板只有23个打席，不能据此确定长期角色。因此，我把牛棚定位为短期校准环境，而不是最终结论。校准期暂定3—4周，每次承担1—2局，重点练习四缝线进区重复性、滑球追打区域，以及变速球形与速度分层。变速球先在确认球形后逐步试到8%—12%，速度差以约8 mph为底线并向9—10 mph恢复；出现右臂疲劳或动作异常即停止加量。",
    ),
    (
        "body",
        "先至少积累60个打席，此后用滚动最近60个打席检查：Zone%不低于46%、首球好球率不低于60%、BB%不高于10%。10%<BB%≤12%时维持负荷，BB%>12%或进区和首球持续未达时回到校准。全部达标后按45、60、75球逐级增加，每级至少两次。60—75球阶段仍健康、过程稳定时回到传统先发；若扩大样本后只在第三轮持续失效，再考虑短先发或跟随投手。这些数值是统一观察尺度的管理规则，不是23个牛棚打席拟合出的最优答案。",
    ),
    ("heading", "三、提交材料与复核"),
    (
        "body",
        "数据截至8月12日，共17场比赛、288个打席。Statcast原始文件有1,241行，其中1行自动坏球没有球种标签，所以球种分析使用1,240个带标签投球。提交前，我重核1,240个投球、80次三振、42次保送和17场比赛，四项与官方记录的差值均为0；同时复算17份原始文件的SHA-256，确保报告数字能追溯到冻结版本。",
    ),
    (
        "body",
        "本周完成两页决策汇总、完整分析底稿、数据来源与指标口径说明、分析迭代记录和复现核查说明。两页版只保留结论、证据、执行门槛与限制，详细分母、置信区间、反例和运行方法放入底稿。MLB-Analytics仓库按原始数据、处理数据、研究记录、脚本、报告和质量检查分层；我最后逐项核对正文、图表和处理表，确认关键数字都能由代码复算。",
    ),
    ("heading", "四、港科大暑研：完成实验与答辩"),
    (
        "body",
        "港科大IEDA暑研本周完成Synthetic-oracle机制实验。我在HKUST HPC4上固定三组随机种子和八条policy pipeline，并把训练与评价分开：每条policy在512个held-out prompts上重新生成6条回答，三组种子共得到73,728条fresh rollout记录。三种子均值中，O-Pro和H-Pro分别在对应信号组内取得最高policy utility；O-Pro的NLL和MSE并非组内最好，但局部regret最低，符合“偏好拟合排序与下游效用排序可能分离”的研究直觉。",
    ),
    (
        "body",
        "我没有把结果写得过满。不同种子仍有波动，实验采用冻结表示层、线性reward head、rank-4 LoRA和单步更新，因此验证的是方法机制，不能直接推广到真人偏好或所有大模型。阶段末，我完成实验汇总、最终汇报和答辩，并提交《Prospective Reward Modeling》第一阶段报告。理论推导、HPC4实验、结果整理和答辩均已完成；真人偏好、更大模型与多轮优化属于后续扩展。",
    ),
    ("heading", "五、阶段产出与总结"),
    (
        "body",
        "这四周让我认识到，分析的价值不只是得到一个数，还要说明数字从哪里来、什么证据支持判断、什么情况会推翻判断，以及下一步如何行动。棒球项目训练了我在不完整公开数据下作出克制决策；港科暑研则把同一思路推进到理论和计算实验。两项任务均已按本阶段范围完成。",
    ),
    ("bullet", "完成今井达也MLB调整方案、两页汇总、完整底稿和可复现项目仓库。"),
    ("bullet", "固定观察门槛、负荷阶梯与退回条件，并保留小样本和跨联盟比较的限制。"),
    ("bullet", "完成港科暑研三随机种子实验、73,728条fresh rollout汇总、最终汇报与答辩。"),
    ("bullet", "确认课程数据分析项目和港科暑研本阶段均已完整完成；后续工作另作扩展。"),
]


def set_run_font(run, *, size: float, bold: bool = False, italic: bool = False) -> None:
    run.font.name = "仿宋_GB2312"
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), "仿宋_GB2312")
    rfonts.set(qn("w:hAnsi"), "仿宋_GB2312")
    rfonts.set(qn("w:eastAsia"), "仿宋_GB2312")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def clear_paragraph(paragraph) -> None:
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)


def apply_role(paragraph, role: str, text: str, templates: dict[str, object]) -> None:
    current_ppr = paragraph._p.pPr
    if current_ppr is not None:
        paragraph._p.remove(current_ppr)
    template_ppr = templates[role]._p.pPr
    if template_ppr is not None:
        paragraph._p.insert(0, deepcopy(template_ppr))
    clear_paragraph(paragraph)
    run = paragraph.add_run(text)
    if role == "note":
        set_run_font(run, size=10.5, italic=True)
    elif role == "heading":
        set_run_font(run, size=12, bold=True)
    else:
        set_run_font(run, size=12)


def ensure_content_paragraphs(doc, count: int) -> list:
    # Title and its following blank paragraph remain untouched. Content begins
    # immediately after the information table, at doc.paragraphs[2].
    while len(doc.paragraphs) - 2 < count:
        doc.add_paragraph()
    while len(doc.paragraphs) - 2 > count:
        paragraph = doc.paragraphs[-1]
        paragraph._p.getparent().remove(paragraph._p)
    return doc.paragraphs[2:]


def build(source: Path, output: Path, content: list[tuple[str, str]]) -> None:
    reference = Document(REFERENCE)
    templates = {
        "note": reference.paragraphs[2],
        "heading": reference.paragraphs[3],
        "body": reference.paragraphs[4],
        "bullet": reference.paragraphs[16],
    }

    doc = Document(source)
    paragraphs = ensure_content_paragraphs(doc, len(content))
    for paragraph, (role, text) in zip(paragraphs, content, strict=True):
        apply_role(paragraph, role, text, templates)

    # Keep the top title centered and make the information-table topic explicit.
    doc.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.core_properties.title = doc.paragraphs[0].text
    doc.core_properties.subject = "本科生短学期每周实习记录"
    doc.core_properties.comments = ""

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main() -> None:
    build(WEEK_DIR / "第3周实习记录.docx", WORK_DIR / "第3周实习记录_候选.docx", WEEK3)
    build(WEEK_DIR / "第4周实习记录.docx", WORK_DIR / "第4周实习记录_候选.docx", WEEK4)
    print(f"week3_chars={sum(len(text) for _, text in WEEK3)}")
    print(f"week4_chars={sum(len(text) for _, text in WEEK4)}")


if __name__ == "__main__":
    main()
