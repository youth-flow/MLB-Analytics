from __future__ import annotations

import argparse
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document


FILES = [
    "蒋总汇总_今井达也MLB调整决策_两页版.docx",
    "底稿01_今井达也MLB调整决策_完整分析.docx",
    "底稿02_MLB数据来源与指标口径.docx",
    "底稿03_MLB分析迭代与反证记录.docx",
    "底稿04_MLB复现与交付核查说明.docx",
]

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
TRACKED = {f"{{{W_NS}}}{name}" for name in ("ins", "del", "moveFrom", "moveTo")}
PLACEHOLDERS = ("<br>", "TODO", "TBD", "XXXX", "〔待补〕", "〔待填〕")


def visible_text(document: Document) -> str:
    chunks = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            chunks.extend(cell.text for cell in row.cells)
    return "\n".join(chunks)


def audit(path: Path) -> dict[str, int]:
    assert path.is_file(), f"missing file: {path}"
    with zipfile.ZipFile(path) as archive:
        assert archive.testzip() is None, f"ZIP CRC failure: {path.name}"
        names = set(archive.namelist())
        assert "word/vbaProject.bin" not in names, f"macro found: {path.name}"
        assert not any(name.startswith("word/embeddings/") for name in names), (
            f"embedded object found: {path.name}"
        )
        assert not any("comments" in name.lower() for name in names), (
            f"comment part found: {path.name}"
        )

        tracked_count = 0
        external_count = 0
        for name in sorted(names):
            if name.endswith(".xml"):
                root = ET.fromstring(archive.read(name))
                tracked_count += sum(1 for node in root.iter() if node.tag in TRACKED)
            elif name.endswith(".rels"):
                root = ET.fromstring(archive.read(name))
                external_count += sum(
                    1
                    for node in root.iter()
                    if node.tag == f"{{{R_NS}}}Relationship"
                    and node.attrib.get("TargetMode") == "External"
                )
        assert tracked_count == 0, f"tracked changes found: {path.name}"
        assert external_count == 0, f"external relationship found: {path.name}"

    document = Document(path)
    assert document.core_properties.author == "杨炎新", f"unexpected author: {path.name}"
    assert document.core_properties.last_modified_by == "杨炎新", (
        f"unexpected last editor: {path.name}"
    )
    assert not document.core_properties.comments, f"generator comment found: {path.name}"
    text = visible_text(document)
    for marker in PLACEHOLDERS:
        assert marker not in text, f"placeholder {marker!r} found: {path.name}"
    assert not re.search(r"_{4,}", text), f"blank underline placeholder found: {path.name}"

    for section in document.sections:
        width_mm = section.page_width.mm
        height_mm = section.page_height.mm
        assert abs(width_mm - 210.0) < 0.6 and abs(height_mm - 297.0) < 0.6, (
            f"non-A4 section in {path.name}: {width_mm:.2f} x {height_mm:.2f} mm"
        )

    return {
        "paragraphs": len(document.paragraphs),
        "tables": len(document.tables),
        "images": len(document.inline_shapes),
        "characters": len(text.replace("\n", "")),
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("root", type=Path)
    args = parser.parse_args()
    for filename in FILES:
        stats = audit(args.root / filename)
        print(
            "PASS|"
            f"{filename}|paragraphs={stats['paragraphs']}|tables={stats['tables']}|"
            f"images={stats['images']}|characters={stats['characters']}"
        )


if __name__ == "__main__":
    main()
