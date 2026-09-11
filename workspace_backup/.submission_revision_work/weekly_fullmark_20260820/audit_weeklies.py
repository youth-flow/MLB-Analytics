from __future__ import annotations

import hashlib
import json
import re
import zipfile
from pathlib import Path

from docx import Document
from lxml import etree
from pypdf import PdfReader


W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
NS = {"w": W, "pr": PKG_REL}


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest().upper()


def east_asia_font(run) -> str | None:
    rpr = run._r.rPr
    if rpr is None or rpr.rFonts is None:
        return None
    return rpr.rFonts.get(f"{{{W}}}eastAsia")


def visible_text(document: Document) -> str:
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def audit_one(path: Path, pdf: Path, week: int) -> dict:
    document = Document(path)
    title = f"第{week}周实习记录"
    assert document.paragraphs[0].text == title
    title_run = document.paragraphs[0].runs[0]
    assert east_asia_font(title_run) == "方正小标宋简体"
    assert round(title_run.font.size.pt, 1) == 16.0
    # The design authority uses the naturally heavy glyphs of 方正小标宋简体;
    # it does not add a separate w:b override.
    assert len(document.sections) == 1
    section = document.sections[0]
    assert round(section.page_width.inches, 2) == 8.27
    assert round(section.page_height.inches, 2) == 11.69
    assert len(document.tables) == 1
    table = document.tables[0]
    assert len(table.rows) == 2 and len(table.columns) == 4
    assert table.cell(0, 1).text.strip() == "杨炎新"
    assert table.cell(0, 3).text.strip() == "3230102355"
    assert table.cell(1, 1).text.strip() == "上海思搏动态智能科技有限公司"
    assert table.cell(1, 3).text.strip().startswith(f"第{week}周：")
    headings = [p for p in document.paragraphs if p.style.name == "Heading 1"]
    assert len(headings) == 5
    for paragraph in document.paragraphs[1:]:
        if not paragraph.text.strip():
            continue
        for run in paragraph.runs:
            if run.text.strip():
                assert east_asia_font(run) == "仿宋_GB2312"
                assert round(run.font.size.pt, 1) in (10.5, 12.0)
    footer_xml = etree.tostring(section.footer._element, encoding="unicode")
    assert "PAGE" in footer_xml
    assert len(PdfReader(str(pdf)).pages) == 2

    with zipfile.ZipFile(path) as archive:
        assert archive.testzip() is None
        names = archive.namelist()
        assert not any(
            name.startswith(("word/comments", "word/embeddings/", "word/activeX/"))
            or "vbaProject" in name
            for name in names
        )
        document_xml = archive.read("word/document.xml")
        settings_xml = archive.read("word/settings.xml")
        document_root = etree.fromstring(document_xml)
        assert not document_root.xpath(".//w:ins | .//w:del", namespaces=NS)
        assert b"trackRevisions" not in settings_xml
        for name in [n for n in names if n.endswith(".rels")]:
            rel_root = etree.fromstring(archive.read(name))
            external = rel_root.xpath(".//pr:Relationship[@TargetMode='External']", namespaces=NS)
            assert not external, f"External relationship in {name}"

    text = visible_text(document)
    assert not re.search(r"TODO|TBD|placeholder|待补|待确认|待完成|作为AI|人工智能生成", text, re.I)
    required = {
        1: ["7月14日", "五个待检验判断", "项目实际推进顺序整理"],
        2: ["1,241行", "1,240个带标签投球", "80次三振", "42次保送", "四项差值为0"],
        3: ["xPV/100为+1.03", "10.14 mph", "8.01 mph", "89球、23个打席", "161球", "实验试跑"],
        4: ["73,728条rollout记录", "17项来源哈希", "课程数据分析项目和港科大暑研第一阶段均已完整完成"],
    }[week]
    for phrase in required:
        assert phrase in text, f"Missing {phrase} in {path.name}"
    return {
        "file": path.name,
        "sha256": sha256(path),
        "pages": 2,
        "headings": 5,
        "characters": len(text.replace("\n", "")),
        "status": "PASS",
    }


def main() -> None:
    root = Path(__file__).resolve().parent
    results = []
    for week in range(1, 5):
        path = root / "output" / f"第{week}周实习记录.docx"
        pdf = root / "render" / "v3" / f"week{week}" / f"第{week}周实习记录.pdf"
        results.append(audit_one(path, pdf, week))
    print(json.dumps({"overall": "PASS", "files": results}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
