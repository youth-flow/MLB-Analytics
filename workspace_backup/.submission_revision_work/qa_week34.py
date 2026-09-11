from pathlib import Path
import re
from zipfile import ZipFile
from lxml import etree
from docx import Document


ROOT = Path(r"C:\Users\yyx\Desktop\大三暑期短学期")
FILES = [
    ROOT / "提交材料" / "01_每周实习记录" / "第3周实习记录.docx",
    ROOT / "提交材料" / "01_每周实习记录" / "第4周实习记录.docx",
]
EXPECTED_NOTE = "说明：以下内容按项目实际推进顺序整理，具体日期以会议材料和实验记录为准。"
FORBIDDEN = ("TODO", "TBD", "{{", "}}", "AI生成", "作为AI", "turn0", "citation")


def cm(value):
    return round(value.cm, 3)


for path in FILES:
    with ZipFile(path) as archive:
        bad = archive.testzip()
        assert bad is None, (path.name, bad)
        etree.fromstring(archive.read("word/document.xml"))
        names = set(archive.namelist())
        assert "word/vbaProject.bin" not in names
        assert not any(name.startswith("word/embeddings/") for name in names)

    doc = Document(path)
    assert len(doc.sections) == 1
    sec = doc.sections[0]
    assert round(sec.page_width.cm, 1) == 21.0
    assert round(sec.page_height.cm, 1) == 29.7
    assert (cm(sec.top_margin), cm(sec.bottom_margin), cm(sec.left_margin), cm(sec.right_margin)) == (
        2.349,
        2.251,
        2.551,
        2.551,
    )

    headings = [p.text.strip() for p in doc.paragraphs if p.style.name == "Heading 1"]
    bullets = [p.text.strip() for p in doc.paragraphs if p.style.name == "List Bullet"]
    note = [p.text.strip() for p in doc.paragraphs if p.text.strip() == EXPECTED_NOTE]
    body_text = "\n".join(p.text for p in doc.paragraphs)
    table_text = "\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
    text = body_text + "\n" + table_text

    assert len(headings) == 5, (path.name, headings)
    assert all(re.match(r"^[一二三四五]、\S", heading) for heading in headings), (path.name, headings)
    assert [heading[0] for heading in headings] == list("一二三四五"), (path.name, headings)
    assert len(bullets) == 4, (path.name, bullets)
    assert len(note) == 1
    assert len(doc.tables) == 1
    assert len(doc.tables[0].rows) == 2 and len(doc.tables[0].columns) == 4
    assert not any(token.lower() in text.lower() for token in FORBIDDEN)
    assert "杨炎新" in text and "3230102355" in text
    assert "上海思搏动态智能科技有限公司" in text

    print(path.name)
    print(f"  paragraphs={len(doc.paragraphs)} headings={len(headings)} bullets={len(bullets)} tables={len(doc.tables)}")
    print(f"  characters={len(text)} crc=PASS xml=PASS layout=PASS placeholders=PASS")
