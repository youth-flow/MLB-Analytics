import argparse
from pathlib import Path

from docx import Document


parser = argparse.ArgumentParser()
parser.add_argument("--root", default=r"C:\Users\yyx\Desktop\大三暑期短学期\提交材料")
parser.add_argument("--contains", default="")
parser.add_argument("--full", action="store_true")
args = parser.parse_args()

root = Path(args.root)
for path in sorted(root.rglob("*.docx")):
    if args.contains and args.contains not in path.name:
        continue
    document = Document(path)
    paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    table_rows = []
    for table in document.tables:
        for row in table.rows:
            table_rows.append(" | ".join(cell.text.strip().replace("\n", " / ") for cell in row.cells))

    character_count = sum(map(len, paragraphs)) + sum(map(len, table_rows))
    print(f"\n### {path.relative_to(root)}")
    print(
        f"paragraphs={len(paragraphs)} tables={len(document.tables)} "
        f"characters={character_count} images={len(document.inline_shapes)}"
    )
    if args.full:
        print("PARAGRAPHS")
        for index, text in enumerate(paragraphs, start=1):
            print(f"P{index}: {text}")
        print("TABLE_ROWS")
        for index, text in enumerate(table_rows, start=1):
            print(f"R{index}: {text}")
    else:
        print("HEAD")
        for index, text in enumerate(paragraphs[:12], start=1):
            print(f"P{index}: {text}")
        print("TAIL")
        start = max(1, len(paragraphs) - 5)
        for index, text in enumerate(paragraphs[-6:], start=start):
            print(f"P{index}: {text}")
