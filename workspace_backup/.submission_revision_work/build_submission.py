from __future__ import annotations

import re
import tempfile
import zipfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Iterable
from xml.etree import ElementTree as ET

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


WORK_ROOT = Path(__file__).resolve().parent
SOURCE_ROOT = WORK_ROOT / "sources"
OUTPUT_ROOT = WORK_ROOT / "output"


@dataclass(frozen=True)
class StyleSpec:
    body_size: float
    line_spacing: float
    margin_top: float
    margin_bottom: float
    margin_left: float
    margin_right: float
    title_size: float
    h1_size: float
    h2_size: float
    h3_size: float
    table_size: float
    caption_size: float
    body_east_asia: str = "宋体"
    body_western: str = "Times New Roman"


STYLE_SPECS = {
    "weekly": StyleSpec(11.5, 1.38, 2.35, 2.25, 2.55, 2.55, 18, 14, 12, 11.5, 10, 9.5),
    "report": StyleSpec(12, 1.5, 2.5, 2.5, 2.7, 2.7, 22, 15, 13, 12, 10, 10),
    # 课程两页正式稿的硬格式：A4、小五（9pt）微软雅黑、首段论点加粗。
    "executive": StyleSpec(
        9, 1.0, 1.7, 1.7, 1.8, 1.8, 17, 12.5, 10.5, 9.5, 9, 9,
        body_east_asia="微软雅黑", body_western="Arial",
    ),
    "draft_long": StyleSpec(11.5, 1.30, 2.25, 2.2, 2.5, 2.5, 18, 14, 12, 11, 9.8, 9.5),
    "draft_short": StyleSpec(11, 1.3, 2.25, 2.15, 2.35, 2.35, 17, 13.5, 11.5, 10.5, 9.5, 9.2),
}


def set_run_font(run, east_asia: str, western: str, size: float, bold: bool = False) -> None:
    run.font.name = western
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    fonts.set(qn("w:eastAsia"), east_asia)
    fonts.set(qn("w:ascii"), western)
    fonts.set(qn("w:hAnsi"), western)
    fonts.set(qn("w:cs"), western)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), "666666")


def set_table_geometry(table, widths_dxa: list[int]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths_dxa[index])


def add_page_number(section, spec: StyleSpec) -> None:
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run()
    set_run_font(run, spec.body_east_asia, spec.body_western, 9)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def configure_document(document: Document, spec: StyleSpec) -> None:
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(spec.margin_top)
    section.bottom_margin = Cm(spec.margin_bottom)
    section.left_margin = Cm(spec.margin_left)
    section.right_margin = Cm(spec.margin_right)
    section.header_distance = Cm(1.2)
    section.footer_distance = Cm(1.2)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = spec.body_western
    normal.font.size = Pt(spec.body_size)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), spec.body_east_asia)
    normal._element.rPr.rFonts.set(qn("w:ascii"), spec.body_western)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), spec.body_western)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Pt(spec.body_size * 2)
    normal.paragraph_format.line_spacing = spec.line_spacing
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.widow_control = True

    heading_east_asia = "微软雅黑" if spec.body_east_asia == "微软雅黑" else "黑体"
    for name, size in (("Heading 1", spec.h1_size), ("Heading 2", spec.h2_size), ("Heading 3", spec.h3_size)):
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), heading_east_asia)
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.paragraph_format.first_line_indent = Pt(0)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True
        style.paragraph_format.line_spacing = 1.1
        if name == "Heading 1":
            style.paragraph_format.space_before = Pt(12)
            style.paragraph_format.space_after = Pt(6)
        elif name == "Heading 2":
            style.paragraph_format.space_before = Pt(9)
            style.paragraph_format.space_after = Pt(4)
        else:
            style.paragraph_format.space_before = Pt(6)
            style.paragraph_format.space_after = Pt(3)

    for section in document.sections:
        add_page_number(section, spec)


def add_inline_text(paragraph, text: str, size: float, east_asia="宋体", western="Times New Roman") -> None:
    text = text.replace("`", "")
    parts = re.split(r"(\*\*.*?\*\*)", text)
    for part in parts:
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**")
        content = part[2:-2] if bold else part
        run = paragraph.add_run(content)
        set_run_font(run, east_asia, western, size, bold=bold)


def add_title(document: Document, metadata: dict[str, str], spec: StyleSpec) -> None:
    kind = metadata["kind"]
    if kind == "report":
        for _ in range(4):
            document.add_paragraph()
        school = document.add_paragraph()
        school.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = school.add_run("浙江大学管理学院")
        set_run_font(run, "黑体", "Arial", 16, bold=True)
        school.paragraph_format.space_after = Pt(22)

        title = document.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title.add_run(metadata["title"])
        set_run_font(run, "黑体", "Arial", spec.title_size, bold=True)
        title.paragraph_format.space_after = Pt(12)

        subtitle = document.add_paragraph()
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = subtitle.add_run(metadata.get("subtitle", ""))
        set_run_font(run, "宋体", "Times New Roman", 14)
        subtitle.paragraph_format.space_after = Pt(56)

        info = [
            ("学生姓名", metadata.get("author", "")),
            ("学号", metadata.get("student_id", "")),
            ("实习单位", metadata.get("company", "")),
            ("校内指导教师", metadata.get("instructor", "")),
            ("实习时间", metadata.get("period", "")),
        ]
        table = document.add_table(rows=len(info), cols=2)
        total = content_width_dxa(document)
        widths = [int(total * 0.26), total - int(total * 0.26)]
        set_table_geometry(table, widths)
        set_table_borders(table)
        for row, (label, value) in zip(table.rows, info):
            row.cells[0].text = label
            row.cells[1].text = value
            for idx, cell in enumerate(row.cells):
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                set_cell_margins(cell, top=120, bottom=120, start=120, end=120)
                if idx == 0:
                    shade_cell(cell, "F2F2F2")
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.first_line_indent = Pt(0)
                    paragraph.paragraph_format.line_spacing = 1.15
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
                    for run in paragraph.runs:
                        set_run_font(run, "黑体" if idx == 0 else "宋体", "Arial" if idx == 0 else "Times New Roman", 11, bold=idx == 0)
        document.add_page_break()
        return

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run(metadata["title"])
    title_east_asia = "微软雅黑" if spec.body_east_asia == "微软雅黑" else "黑体"
    set_run_font(run, title_east_asia, "Arial", spec.title_size, bold=True)

    subtitle_text = metadata.get("subtitle", "")
    if subtitle_text:
        subtitle = document.add_paragraph()
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle.paragraph_format.first_line_indent = Pt(0)
        subtitle.paragraph_format.space_after = Pt(6)
        run = subtitle.add_run(subtitle_text)
        set_run_font(run, spec.body_east_asia, spec.body_western, spec.body_size)

    if kind == "weekly":
        info = [
            ["学生", metadata.get("author", ""), "学号", metadata.get("student_id", "")],
            ["实习单位", metadata.get("company", ""), "记录主题", metadata.get("period", "")],
        ]
        add_table(document, info, spec, header=False, preferred_weights=[1.0, 1.6, 1.0, 1.8])
    else:
        meta_line = metadata.get("meta_line", "")
        if meta_line:
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Pt(0)
            paragraph.paragraph_format.space_after = Pt(8)
            add_inline_text(
                paragraph,
                meta_line,
                spec.caption_size,
                east_asia=spec.body_east_asia,
                western=spec.body_western,
            )


def content_width_dxa(document: Document) -> int:
    section = document.sections[0]
    # python-docx exposes page geometry in EMU; Word table widths use twentieths
    # of a point (DXA), and 1 DXA = 635 EMU.
    return int((section.page_width - section.left_margin - section.right_margin) / 635)


def table_weights(rows: list[list[str]], preferred_weights: list[float] | None = None) -> list[float]:
    ncols = max(len(row) for row in rows)
    if preferred_weights and len(preferred_weights) == ncols:
        return preferred_weights
    scores = []
    for col in range(ncols):
        lengths = [max(2, min(40, len(row[col]) if col < len(row) else 0)) for row in rows]
        score = max(4.0, sum(lengths) / max(1, len(lengths)))
        scores.append(score)
    if ncols >= 3:
        scores = [min(score, 18.0) for score in scores]
    return scores


def add_table(
    document: Document,
    rows: list[list[str]],
    spec: StyleSpec,
    header: bool = True,
    preferred_weights: list[float] | None = None,
) -> None:
    if not rows:
        return
    ncols = max(len(row) for row in rows)
    normalized = [row + [""] * (ncols - len(row)) for row in rows]
    table = document.add_table(rows=len(normalized), cols=ncols)
    total = content_width_dxa(document)
    weights = table_weights(normalized, preferred_weights)
    weight_sum = sum(weights)
    widths = [max(600, int(total * weight / weight_sum)) for weight in weights]
    widths[-1] += total - sum(widths)
    set_table_geometry(table, widths)
    set_table_borders(table)
    if header:
        set_repeat_table_header(table.rows[0])

    for rindex, (row, values) in enumerate(zip(table.rows, normalized)):
        prevent_row_split(row)
        for cindex, (cell, value) in enumerate(zip(row.cells, values)):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if header and rindex == 0:
                shade_cell(cell, "EDEDED")
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.first_line_indent = Pt(0)
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.15
            if header and rindex == 0:
                # Do not leave a repeated header row alone at the bottom of a page.
                paragraph.paragraph_format.keep_with_next = True
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if (header and rindex == 0) or len(value) <= 12 else WD_ALIGN_PARAGRAPH.LEFT
            paragraph.clear()
            add_inline_text(
                paragraph,
                value,
                spec.table_size,
                east_asia=("微软雅黑" if spec.body_east_asia == "微软雅黑" else "黑体")
                if header and rindex == 0
                else spec.body_east_asia,
                western="Arial" if header and rindex == 0 else spec.body_western,
            )
            if header and rindex == 0:
                for run in paragraph.runs:
                    run.font.bold = True
    spacer = document.add_paragraph()
    spacer.paragraph_format.first_line_indent = Pt(0)
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing = 0.3


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        parts = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", cell or "---") for cell in parts):
            rows.append(parts)
        index += 1
    return rows, index


def add_body_paragraph(document: Document, text: str, spec: StyleSpec, lead=False) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Pt(spec.body_size * 2)
    paragraph.paragraph_format.line_spacing = spec.line_spacing
    paragraph.paragraph_format.space_after = Pt(0)
    if lead:
        paragraph.paragraph_format.left_indent = Pt(spec.body_size)
        paragraph.paragraph_format.right_indent = Pt(spec.body_size)
        paragraph.paragraph_format.space_before = Pt(3)
        paragraph.paragraph_format.space_after = Pt(5)
    add_inline_text(
        paragraph,
        text,
        spec.body_size,
        east_asia=spec.body_east_asia,
        western=spec.body_western,
    )
    if lead:
        for run in paragraph.runs:
            run.font.bold = True


def add_reference_paragraph(document: Document, text: str, spec: StyleSpec) -> None:
    """Format one bibliography entry without stretched justified spacing."""
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.left_indent = Pt(spec.body_size * 2)
    paragraph.paragraph_format.first_line_indent = Pt(-spec.body_size * 2)
    paragraph.paragraph_format.line_spacing = 1.2
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(3)
    add_inline_text(
        paragraph,
        text,
        spec.body_size,
        east_asia=spec.body_east_asia,
        western=spec.body_western,
    )


def add_list_item(document: Document, text: str, spec: StyleSpec, numbered=False) -> None:
    style_name = "List Number" if numbered else "List Bullet"
    paragraph = document.add_paragraph(style=style_name)
    paragraph.paragraph_format.left_indent = Cm(0.75)
    paragraph.paragraph_format.first_line_indent = Cm(-0.35)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(1.5)
    paragraph.paragraph_format.line_spacing = max(1.15, spec.line_spacing - 0.1)
    add_inline_text(
        paragraph,
        text,
        spec.body_size,
        east_asia=spec.body_east_asia,
        western=spec.body_western,
    )


def add_image(document: Document, path: Path, caption: str, width_cm: float, spec: StyleSpec) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    run.add_picture(str(path), width=Cm(width_cm))
    caption_paragraph = document.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.paragraph_format.first_line_indent = Pt(0)
    caption_paragraph.paragraph_format.space_before = Pt(3)
    caption_paragraph.paragraph_format.space_after = Pt(5)
    add_inline_text(
        caption_paragraph,
        caption,
        spec.caption_size,
        east_asia=spec.body_east_asia,
        western=spec.body_western,
    )


def normalize_package_metadata(path: Path) -> None:
    """Remove template/generator traces and stale package statistics."""
    app_ns = "http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"
    with zipfile.ZipFile(path, "r") as source_zip:
        entries = {item.filename: source_zip.read(item.filename) for item in source_zip.infolist()}

    app_name = "docProps/app.xml"
    if app_name in entries:
        root = ET.fromstring(entries[app_name])
        for tag in (
            "Pages",
            "Words",
            "Characters",
            "CharactersWithSpaces",
            "Lines",
            "Paragraphs",
        ):
            node = root.find(f"{{{app_ns}}}{tag}")
            if node is not None:
                root.remove(node)
        application = root.find(f"{{{app_ns}}}Application")
        if application is None:
            application = ET.SubElement(root, f"{{{app_ns}}}Application")
        application.text = "Office Open XML"
        version = root.find(f"{{{app_ns}}}AppVersion")
        if version is not None:
            root.remove(version)
        entries[app_name] = ET.tostring(root, encoding="utf-8", xml_declaration=True)

    with tempfile.NamedTemporaryFile(delete=False, suffix=".docx", dir=path.parent) as handle:
        temp_path = Path(handle.name)
    try:
        with zipfile.ZipFile(temp_path, "w", compression=zipfile.ZIP_DEFLATED) as target_zip:
            for name in sorted(entries):
                target_zip.writestr(name, entries[name])
        temp_path.replace(path)
    finally:
        if temp_path.exists():
            temp_path.unlink()


def parse_source(path: Path) -> tuple[dict[str, str], list[str]]:
    text = path.read_text(encoding="utf-8")
    parts = text.split("---", 2)
    if len(parts) != 3:
        raise ValueError(f"Missing front matter in {path}")
    metadata: dict[str, str] = {}
    for line in parts[1].splitlines():
        if not line.strip() or line.lstrip().startswith("#"):
            continue
        key, value = line.split(":", 1)
        metadata[key.strip()] = value.strip()
    return metadata, parts[2].strip().splitlines()


def build_document(source_path: Path) -> Path:
    metadata, lines = parse_source(source_path)
    kind = metadata["kind"]
    spec = STYLE_SPECS[kind]
    document = Document()
    configure_document(document, spec)
    add_title(document, metadata, spec)

    index = 0
    pending_widths: list[float] | None = None
    current_heading = ""
    while index < len(lines):
        stripped = lines[index].strip()
        if not stripped:
            index += 1
            continue
        if stripped == "[[PAGEBREAK]]":
            document.add_page_break()
            index += 1
            continue
        if stripped.startswith("[[COLWEIGHTS:"):
            payload = stripped[len("[[COLWEIGHTS:") : -2]
            pending_widths = [float(item.strip()) for item in payload.split(",")]
            index += 1
            continue
        if stripped.startswith("[[IMAGE:"):
            payload = stripped[len("[[IMAGE:") : -2]
            image_path, caption, width = payload.split("|", 2)
            add_image(document, Path(image_path), caption, float(width), spec)
            index += 1
            continue
        if stripped.startswith("[[EQUATION:"):
            payload = stripped[len("[[EQUATION:") : -2]
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Pt(0)
            paragraph.paragraph_format.space_before = Pt(4)
            paragraph.paragraph_format.space_after = Pt(4)
            add_inline_text(paragraph, payload, spec.body_size, east_asia="宋体", western="Cambria Math")
            index += 1
            continue
        if stripped.startswith("|"):
            rows, index = parse_table(lines, index)
            add_table(document, rows, spec, header=True, preferred_weights=pending_widths)
            pending_widths = None
            continue
        if stripped.startswith("### "):
            paragraph = document.add_paragraph(stripped[4:], style="Heading 2")
            index += 1
            continue
        if stripped.startswith("#### "):
            paragraph = document.add_paragraph(stripped[5:], style="Heading 3")
            index += 1
            continue
        if stripped.startswith("## "):
            current_heading = stripped[3:]
            paragraph = document.add_paragraph(current_heading, style="Heading 1")
            index += 1
            continue
        if stripped.startswith("> "):
            add_body_paragraph(document, stripped[2:], spec, lead=True)
            index += 1
            continue
        if kind == "report" and current_heading == "参考文献":
            add_reference_paragraph(document, stripped, spec)
            index += 1
            continue
        if stripped.startswith("- "):
            add_list_item(document, stripped[2:], spec, numbered=False)
            index += 1
            continue
        numbered = re.match(r"^\d+\.\s+(.*)$", stripped)
        if numbered:
            add_list_item(document, numbered.group(1), spec, numbered=True)
            index += 1
            continue
        add_body_paragraph(document, stripped, spec)
        index += 1

    for section in document.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(spec.margin_top)
        section.bottom_margin = Cm(spec.margin_bottom)
        section.left_margin = Cm(spec.margin_left)
        section.right_margin = Cm(spec.margin_right)
    properties = document.core_properties
    properties.author = metadata.get("author", "杨炎新")
    properties.last_modified_by = metadata.get("author", "杨炎新")
    properties.title = metadata["title"]
    properties.subject = ""
    properties.comments = ""
    properties.keywords = ""
    properties.category = ""
    properties.identifier = ""
    properties.language = "zh-CN"
    properties.version = "1.0"
    properties.revision = 1
    generated_at = datetime.now().replace(microsecond=0)
    properties.created = generated_at
    properties.modified = generated_at
    OUTPUT_ROOT.mkdir(parents=True, exist_ok=True)
    output_path = OUTPUT_ROOT / metadata["filename"]
    document.save(output_path)
    normalize_package_metadata(output_path)
    return output_path


def main() -> None:
    outputs = [build_document(path) for path in sorted(SOURCE_ROOT.glob("*.md"))]
    for output in outputs:
        print(output)


if __name__ == "__main__":
    main()
