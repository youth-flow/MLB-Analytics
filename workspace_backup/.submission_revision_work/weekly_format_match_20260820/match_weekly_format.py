from __future__ import annotations

import copy
import hashlib
import json
import os
import re
import zipfile
from pathlib import Path

from docx import Document
from lxml import etree


W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W}


def qn(local: str) -> str:
    return f"{{{W}}}{local}"


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest().upper()


def set_paragraph_text(paragraph: etree._Element, text: str) -> None:
    """Replace visible paragraph text while preserving pPr and the first run's rPr."""
    source_runs = paragraph.xpath("./w:r", namespaces=NS)
    if not source_runs:
        run = etree.Element(qn("r"))
    else:
        run = copy.deepcopy(source_runs[0])

    for child in list(run):
        if child.tag != qn("rPr"):
            run.remove(child)
    text_node = etree.SubElement(run, qn("t"))
    if text[:1].isspace() or text[-1:].isspace():
        text_node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    text_node.text = text

    for child in list(paragraph):
        if child.tag != qn("pPr"):
            paragraph.remove(child)
    paragraph.append(run)


def extract_target(path: Path) -> dict:
    document = Document(path)
    paragraphs = []
    started = False
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not started:
            if paragraph.style.name == "Heading 1" and text:
                started = True
            else:
                continue
        if not text:
            continue
        if paragraph.style.name == "Heading 1":
            role = "heading"
        elif paragraph.style.name == "List Bullet":
            role = "bullet"
        else:
            role = "body"
        paragraphs.append({"role": role, "text": text})

    heading_count = sum(item["role"] == "heading" for item in paragraphs)
    if heading_count != 5:
        raise ValueError(f"{path.name}: expected 5 headings, found {heading_count}")

    match = re.match(r"第([234])周", path.name)
    if not match:
        raise ValueError(f"Cannot infer week number from {path.name}")

    # The fourth-week source ends with two consecutive body paragraphs whose
    # second paragraph is only one short sentence.  Under the larger first-week
    # typography that sentence alone creates an almost-empty third page.  Join
    # the two paragraphs without changing any visible wording; this is a
    # pagination-only adjustment and preserves the complete source text.
    if match.group(1) == "4" and len(paragraphs) >= 2:
        if paragraphs[-2]["role"] == "body" and paragraphs[-1]["role"] == "body":
            paragraphs[-2]["text"] += paragraphs[-1]["text"]
            paragraphs.pop()

    return {
        "source": str(path),
        "week": match.group(1),
        "title": f"第{match.group(1)}周实习记录",
        "topic": document.tables[0].cell(1, 3).text.strip(),
        "paragraphs": paragraphs,
    }


def build_one(reference: Path, target: Path, output: Path) -> dict:
    target_data = extract_target(target)
    reference_document = Document(reference)
    heading_style_id = reference_document.styles["Heading 1"].style_id
    bullet_style_id = reference_document.styles["List Bullet"].style_id
    with zipfile.ZipFile(reference, "r") as source_zip:
        document_xml = source_zip.read("word/document.xml")
        root = etree.fromstring(document_xml)
        body = root.find("w:body", NS)
        if body is None:
            raise ValueError("Reference document has no w:body")

        children = list(body)
        title_template = next(child for child in children if child.tag == qn("p"))
        table_template = next(child for child in children if child.tag == qn("tbl"))
        table_index = children.index(table_template)
        spacer_template = next(
            child
            for child in children[table_index + 1 :]
            if child.tag == qn("p") and not child.xpath(".//w:t", namespaces=NS)
        )
        heading_template = next(
            child
            for child in children
            if child.tag == qn("p")
            and child.xpath(
                f"./w:pPr/w:pStyle[@w:val='{heading_style_id}']", namespaces=NS
            )
        )
        body_template = next(
            child
            for child in children[children.index(heading_template) + 1 :]
            if child.tag == qn("p")
            and not child.xpath(
                f"./w:pPr/w:pStyle[@w:val='{heading_style_id}']", namespaces=NS
            )
            and not child.xpath(
                f"./w:pPr/w:pStyle[@w:val='{bullet_style_id}']", namespaces=NS
            )
            and child.xpath(".//w:t", namespaces=NS)
        )
        bullet_template = next(
            child
            for child in children
            if child.tag == qn("p")
            and child.xpath(
                f"./w:pPr/w:pStyle[@w:val='{bullet_style_id}']", namespaces=NS
            )
        )
        sect_pr = next(child for child in children if child.tag == qn("sectPr"))

        title = copy.deepcopy(title_template)
        set_paragraph_text(title, target_data["title"])
        table = copy.deepcopy(table_template)
        topic_cell = table.xpath("./w:tr[2]/w:tc[4]", namespaces=NS)[0]
        topic_paragraph = topic_cell.xpath("./w:p[1]", namespaces=NS)[0]
        set_paragraph_text(topic_paragraph, target_data["topic"])
        spacer = copy.deepcopy(spacer_template)

        rebuilt = [title, table, spacer]
        template_by_role = {
            "heading": heading_template,
            "body": body_template,
            "bullet": bullet_template,
        }
        for item in target_data["paragraphs"]:
            paragraph = copy.deepcopy(template_by_role[item["role"]])
            set_paragraph_text(paragraph, item["text"])
            rebuilt.append(paragraph)
        rebuilt.append(copy.deepcopy(sect_pr))

        for child in list(body):
            body.remove(child)
        for child in rebuilt:
            body.append(child)

        output.parent.mkdir(parents=True, exist_ok=True)
        new_document_xml = etree.tostring(
            root, xml_declaration=True, encoding="UTF-8", standalone="yes"
        )
        with zipfile.ZipFile(output, "w") as destination_zip:
            for info in source_zip.infolist():
                payload = (
                    new_document_xml
                    if info.filename == "word/document.xml"
                    else source_zip.read(info.filename)
                )
                destination_zip.writestr(info, payload)

    result_document = Document(output)
    output_body = []
    started = False
    for paragraph in result_document.paragraphs:
        text = paragraph.text.strip()
        if not started and paragraph.style.name == "Heading 1" and text:
            started = True
        if started and text:
            role = (
                "heading"
                if paragraph.style.name == "Heading 1"
                else "bullet"
                if paragraph.style.name == "List Bullet"
                else "body"
            )
            output_body.append({"role": role, "text": text})

    if output_body != target_data["paragraphs"]:
        raise AssertionError(f"Text/style mismatch after building {output.name}")
    if result_document.paragraphs[0].text != target_data["title"]:
        raise AssertionError(f"Title mismatch in {output.name}")
    if result_document.tables[0].cell(1, 3).text.strip() != target_data["topic"]:
        raise AssertionError(f"Topic mismatch in {output.name}")

    reference_parts = {}
    output_parts = {}
    with zipfile.ZipFile(reference) as z:
        reference_parts = {
            info.filename: hashlib.sha256(z.read(info.filename)).hexdigest()
            for info in z.infolist()
        }
    with zipfile.ZipFile(output) as z:
        bad = z.testzip()
        if bad is not None:
            raise AssertionError(f"CRC failure in {output.name}: {bad}")
        output_parts = {
            info.filename: hashlib.sha256(z.read(info.filename)).hexdigest()
            for info in z.infolist()
        }
    changed_parts = [
        name for name in reference_parts if reference_parts[name] != output_parts.get(name)
    ]
    if changed_parts != ["word/document.xml"]:
        raise AssertionError(
            f"Unexpected package changes in {output.name}: {changed_parts}"
        )

    return {
        **target_data,
        "output": str(output),
        "sha256": sha256(output),
        "changed_parts": changed_parts,
        "paragraph_count": len(output_body),
    }


def main() -> None:
    task_root = Path(os.environ["TASK_ROOT"])
    reference = Path(os.environ["REFERENCE_DOCX"])
    target_dir = Path(os.environ["TARGET_DIR"])
    expected_reference_hash = (
        "07B267887710B0EF542D43EA1D6556CD4DA4D4BE6A83C9DFA967367F44BE0254"
    )
    if sha256(reference) != expected_reference_hash:
        raise AssertionError("Reference hash changed; fresh distillation is required")

    targets = [
        target_dir / "第2周实习记录_数据治理与理论推导.docx",
        target_dir / "第3周实习记录_分层分析与HPC4实验.docx",
        target_dir / "第4周实习记录_决策交付与暑研答辩.docx",
    ]
    results = []
    for target in targets:
        if not target.exists():
            raise FileNotFoundError(target)
        results.append(build_one(reference, target, task_root / "output" / target.name))

    if sha256(reference) != expected_reference_hash:
        raise AssertionError("Reference was modified during authoring")
    report = {
        "reference": str(reference),
        "reference_sha256": sha256(reference),
        "outputs": results,
    }
    (task_root / "build_report.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps(report, ensure_ascii=True, indent=2))


if __name__ == "__main__":
    main()
