from __future__ import annotations

import csv
import argparse
import re
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

WORK = Path(__file__).resolve().parent
ROOT = WORK.parent.parent
parser = argparse.ArgumentParser(description='Build the final two submission documents.')
parser.add_argument('--project-root', type=Path)
parser.add_argument('--output-dir', type=Path)
ARGS = parser.parse_args()
PROJECT = ARGS.project_root or next((p for p in WORK.parents if (p/'data/processed/mlb_process_metrics_by_role.csv').is_file()), ROOT/'MLB-Analytics')
ASSETS = WORK / 'assets'
OUT = ARGS.output_dir or (PROJECT/'final_submission' if PROJECT in WORK.parents else ROOT/'提交材料'/'最终提交两份')
OUT.mkdir(parents=True, exist_ok=True)
BLUE = '#355B74'
GRAY = '#979797'


def run_font(run, size=12, bold=False, font='宋体'):
    run.font.name = 'Times New Roman' if font == '宋体' else font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), font)


def configure(doc, size=12, executive=False):
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.footer_distance = Cm(0.85)
    font = '微软雅黑' if executive else '宋体'
    normal = doc.styles['Normal']
    normal.font.name = 'Times New Roman' if not executive else font
    normal.font.size = Pt(size)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), font)
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    doc.styles['Subtitle'].font.italic = False
    doc.styles['Subtitle'].font.color.rgb = RGBColor(0, 0, 0)
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(14 if executive else 18.5)
    pf.space_after = Pt(5 if executive else 7)
    pf.first_line_indent = Pt(0 if executive else size * 2)
    pf.widow_control = True
    for name, pt in [('Title', 17 if executive else 21), ('Heading 1', 11 if executive else 15), ('Heading 2', 10 if executive else 12.5)]:
        style = doc.styles[name]
        style.font.name = '微软雅黑' if executive else '黑体'
        style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '微软雅黑' if executive else '黑体')
        style.font.size = Pt(pt)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.first_line_indent = Pt(0)
        style.paragraph_format.space_before = Pt(7)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.line_spacing = 1.15
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.first_line_indent = Pt(0)
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    props = doc.core_properties
    props.author = '杨炎新'
    props.last_modified_by = '杨炎新'
    props.comments = ''
    props.subject = ''
    props.keywords = ''
    props.language = 'zh-CN'
    settings = doc.settings.element
    for node in list(settings.findall(qn('w:doNotAutoHyphenate'))):
        settings.remove(node)


def ptext(doc, text, size=12, executive=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for token in re.split(r'(\*\*.*?\*\*)', text):
        if not token:
            continue
        bold = token.startswith('**') and token.endswith('**')
        run_font(p.add_run(token[2:-2] if bold else token.replace('`', '')), size, bold, '微软雅黑' if executive else '宋体')
    return p


def caption(doc, text, executive=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    p.paragraph_format.space_after = Pt(7)
    run_font(p.add_run(text), 8.5 if executive else 9.5, font='微软雅黑' if executive else '宋体')


def image_row(doc, paths, height, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    for i, path in enumerate(paths):
        if i:
            p.add_run('   ')
        shape = p.add_run().add_picture(str(path), height=Cm(height))
        shape._inline.docPr.set('descr', text)
    caption(doc, text)


def figure(doc, name, width, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    shape = p.add_run().add_picture(str(ASSETS / name), width=Cm(width))
    shape._inline.docPr.set('descr', text)
    caption(doc, text)


def table(doc, rows, widths, size=10.5, executive=False):
    t = doc.add_table(rows=0, cols=len(widths))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    borders = OxmlElement('w:tblBorders')
    for side in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        edge = OxmlElement('w:' + side)
        for k, v in {'val':'single', 'sz':'4', 'color':'D9D9D9'}.items():
            edge.set(qn('w:' + k), v)
        borders.append(edge)
    t._tbl.tblPr.append(borders)
    for col, width in zip(t.columns, widths):
        col.width = Cm(width)
    for idx, row in enumerate(rows):
        cells = t.add_row().cells
        trPr = cells[0]._tc.getparent().get_or_add_trPr()
        trPr.append(OxmlElement('w:cantSplit'))
        if idx == 0:
            trPr.append(OxmlElement('w:tblHeader'))
        for cell, value, width in zip(cells, row, widths):
            cell.width = Cm(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcPr = cell._tc.get_or_add_tcPr()
            mar = OxmlElement('w:tcMar')
            for side, val in [('top','75'),('bottom','75'),('left','90'),('right','90')]:
                node=OxmlElement('w:'+side);node.set(qn('w:w'),val);node.set(qn('w:type'),'dxa');mar.append(node)
            tcPr.append(mar)
            if idx == 0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EEEEEE');tcPr.append(shade)
            p = cell.paragraphs[0]
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.keep_with_next = idx == 0
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx == 0 or len(str(value)) < 18 else WD_ALIGN_PARAGRAPH.LEFT
            run_font(p.add_run(str(value)), size, idx == 0, '微软雅黑' if executive else '宋体')
    after = doc.add_paragraph()
    after.paragraph_format.first_line_indent=Pt(0)
    after.paragraph_format.space_after=Pt(0)
    after.paragraph_format.line_spacing=Pt(3)
    run_font(after.add_run(''),3)
    return t


def math_run(text):
    r=OxmlElement('m:r'); t=OxmlElement('m:t');t.text=text;r.append(t);return r


def local_regret(doc):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.line_spacing=1.1
    m=OxmlElement('m:oMath')
    m.append(math_run('Reg'))
    m.append(math_run('(r) = U(r*) − U(r)'))
    p._p.append(m)


def charts():
    f=lambda n:ImageFont.truetype('C:/Windows/Fonts/msyh.ttc', n)
    # Fresh figures, plotted directly from frozen results; no photo manipulation.
    im=Image.new('RGB',(1440,450),'white'); d=ImageDraw.Draw(im)
    groups=[('三振率 K%',27.8,80/288*100),('保送率 BB%',7.0,42/288*100)]
    for i,(label,npb,mlb) in enumerate(groups):
        y=70+i*150;d.text((12,y+27),label,font=f(31),fill='black')
        for j,(value,color) in enumerate([(npb,GRAY),(mlb,BLUE)]):
            yy=y+j*51;end=350+value/35*830
            d.rectangle((350,yy,end,yy+34),fill=color)
            d.text((end+16,yy-3),f'{value:.1f}%',font=f(28),fill=color)
    d.rectangle((410,378,447,403),fill=GRAY);d.text((460,374),'NPB 2025',font=f(27),fill='black')
    d.rectangle((770,378,807,403),fill=BLUE);d.text((820,374),'MLB 2026 冻结窗口',font=f(27),fill='black')
    im.save(ASSETS/'control_comparison.png')
    rows=list(csv.DictReader((PROJECT/'data/processed/mlb_process_metrics_by_role.csv').open(encoding='utf-8-sig')))
    im=Image.new('RGB',(1440,660),'white');d=ImageDraw.Draw(im)
    x=lambda v:280+float(v)*9.8
    for tick in range(0,101,20):
        xx=x(tick);d.line((xx,45,xx,520),fill='#DDDDDD',width=2);d.text((xx-24,535),str(tick)+'%',font=f(25),fill='#555555')
    for i,(label,key) in enumerate([('进区率','zone'),('首球好球率','first_pitch_strike'),('挥空率','whiff'),('追打率','chase')]):
        y=75+i*120;d.text((12,y+4),label,font=f(31),fill='black')
        for j,(row,color) in enumerate(zip(rows,[GRAY,BLUE])):
            yy=y+j*34;v=float(row[key+'_pct']);lo=float(row[key+'_ci_low']);hi=float(row[key+'_ci_high'])
            d.line((x(lo),yy,x(hi),yy),fill=color,width=8)
            d.ellipse((x(v)-9,yy-9,x(v)+9,yy+9),fill=color)
            d.text((1290,yy-19),f'{v:.1f}%',font=f(27),fill=color)
    d.text((300,599),'灰色 先发265打席    蓝色 牛棚23打席    线段为Wilson 95%区间',font=f(27),fill='black')
    im.save(ASSETS/'role_intervals.png')


DATA_TABLE=[['来源','主要用途'],['MLB StatsAPI','比赛日志、角色、出局数、三振及保送'],['Baseball Savant','逐球、球种、挥棒、xERA与官方球种价值'],['NPB官方及NPB Basement','职业历史与NPB球种结构参照'],['MLB.com报道','补充角色改变的时间背景']]
HYPOTHESIS_TABLE=[['待检验解释','重点检查','需要保留的相反证据'],['球威普遍下降','三振、挥空、四缝线球速','三项未同步走低，则不先归因于球威'],['球数控制不足','进区、首球好球、保送','若过程改善而失分不降，另查接触质量'],['球种配套不足','滑球价值、第三球种速度层','高使用率本身不代表应削减核心球种'],['侧别或任务长度影响','左右打、角色、打线轮次','差异不稳或区间宽，不作永久角色判断']]
REPRO_TABLE=[['保留内容','复核时的用途'],['data/raw/ 与来源清单','确认当时下载了什么、来自哪里、是否改变'],['data/processed/ 分组表','核对分子分母、球种和角色样本'],['scripts/ 与 tests/','按固定口径重算，并检查关键统计关系'],['qa/ 对账记录','确认比赛数、投球数、三振及保送闭合']]
METRIC_TABLE=[['指标','本窗口分子与分母','结果'],['三振率','80次三振 ÷ 288打席','27.8%'],['保送率','42次保送 ÷ 288打席','14.6%'],['进区率','541个区内球 ÷ 1,240球','43.6%'],['首球好球率','149个好球结果 ÷ 288个打席首球','51.7%'],['挥空率','175次挥空 ÷ 542次挥棒','32.3%'],['追打率','190次区外挥棒 ÷ 699个区外球','27.2%']]
PITCH_TABLE=[['四缝线分组','投球数','挥空率','投手RV/100'],['对左打','311','15.4%','−1.27'],['对右打','229','27.0%','+1.47']]
ACTION_TABLE=[['步骤','观察标准','处理方式'],['初始校准','3—4周多局牛棚；先完成至少60打席','优先检查健康、进区、首球与对左打配套'],['过程观察','滚动60打席：进区≥46%、首球≥60%、保送≤10%','任一缺项不升级；保送>12%或前两项同时缺项，回校准'],['负荷增加','45、60、75球；每级至少两次','条件未满足维持；负荷相关退化退一级；健康异常先复查'],['角色评估','完成75球阶段，各级至少两次且过程稳定','再评估传统先发或短先发安排']]
CONFIG_TABLE=[['实验环节','固定设置'],['奖励模型','冻结1,536维表示，仅训练无偏置线性头'],['策略空间','最后一层q_proj与v_proj的rank-4 LoRA-B'],['策略更新','β＝0.2；单步自然梯度'],['比较信号','O为精确合成奖励差；H为重复模拟偏好标签'],['评价样本','每策略512个正式测试提示，各重新生成6个回答']]
RESULT_TABLE=[['方法','偏好NLL 越低越好','策略效用J 越高越好'],['O-MLE','0.68345 ± 0.00109','0.01599 ± 0.00960'],['O-Pro','0.73285 ± 0.00318','0.02068 ± 0.01837'],['H-MLE','0.68905 ± 0.00217','0.01768 ± 0.00817'],['H-Pro','0.75072 ± 0.00586','0.02000 ± 0.01729']]


def blocks(text):
    return [p.strip() for p in re.split(r'\n\s*\n',text.strip()) if p.strip()]


def sections(path):
    parts=re.split(r'^## ',path.read_text(encoding='utf-8'),flags=re.M)
    out=[]
    for part in parts[1:]:
        title,_,body=part.partition('\n');out.append((title.strip(),body.strip()))
    return out


def report():
    frame=sections(WORK/'report_frame.md')
    mlb=sections(WORK/'mlb_prose.md')
    assert len(frame)==7 and len(mlb)==8, (len(frame),len(mlb))
    pages=[frame[0]]+mlb+frame[1:]
    mlb_headings=['研究问题与分析假设','数据来源与冻结范围','数据清洗与指标口径','控球表现与调整优先级','球种结构与打者侧别','先发牛棚与打线轮次','调整步骤与评估条件','复现整理与判断修订']
    doc=Document();configure(doc)
    doc.core_properties.title='本科生短学期实习报告'
    for page,(heading,body) in enumerate(pages,1):
        if 2 <= page <= 9:
            heading=mlb_headings[page-2]
        if page==1:
            p=doc.add_paragraph('本科生短学期实习报告',style='Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p=doc.add_paragraph('MLB投手数据分析与港科大IEDA暑研',style='Subtitle')
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0)
            for r in p.runs:run_font(r,12)
            for line in ['姓名  杨炎新     学号  3230102355','接收单位  上海思搏动态智能科技有限公司','指导教师  陈发动     实践时间  2026年暑期']:
                p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_after=Pt(3)
                run_font(p.add_run(line),10.5)
        h=doc.add_paragraph(heading,style='Heading 1')
        if page>1:h.paragraph_format.page_break_before=True
        for idx,chunk in enumerate(blocks(body)):
            if chunk=='[[PPT]]':
                image_row(doc,[ASSETS/'阶段汇报总览.png',ASSETS/'基础手写.jpg'],5.7,'图3  第一周阶段汇报总览与基础学习手稿')
            elif chunk=='[[NOTES_A]]':
                image_row(doc,[ASSETS/f'theory-notes-0{i}.jpg' for i in (1,2,3)],6.2,'图4  第二周Prospective Reward Modeling早期推导笔记')
                local_regret(doc)
            elif chunk=='[[NOTES_B]]':
                table(doc,CONFIG_TABLE,[3.0,13.6]);caption(doc,'表7  暑研机制实验的主要控制条件')
            elif chunk=='[[RESULTS]]':
                table(doc,RESULT_TABLE,[2.4,7.1,7.1]);caption(doc,'表8  三个随机种子的均值与样本标准差')
            elif chunk=='[[DEFENSE]]':
                figure(doc,'答辩现场.jpg',15.8,'图5  2026年8月港科大IEDA暑研最终答辩现场')
            elif chunk.startswith('### '):
                doc.add_paragraph(chunk[4:],style='Heading 2')
            else:
                p=ptext(doc,chunk)
                if page==15 and (chunk.startswith('[') or chunk.startswith('项目与')):
                    p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.line_spacing=1.1;p.paragraph_format.space_after=Pt(4)
                    for r in p.runs:run_font(r,9)
            if page==3 and idx==1:
                table(doc,DATA_TABLE,[4.4,12.2]);caption(doc,'表2  主要来源与分析用途')
            if page==2 and idx==1:
                table(doc,HYPOTHESIS_TABLE,[3.2,4.7,8.7]);caption(doc,'表1  研究假设与核查方法')
            if page==4 and idx==1:
                table(doc,METRIC_TABLE,[3.4,10.3,2.9]);caption(doc,'表3  主要过程指标及实际分母')
            if page==5 and idx==1:
                figure(doc,'control_comparison.png',15.6,'图1  三振率保持接近而保送率上升  跨联盟数值仅作历史参照')
            if page==6 and idx==3:
                table(doc,PITCH_TABLE,[4.8,3.1,4.1,4.6]);caption(doc,'表4  四缝线面对不同侧别打者的样本内结果')
            if page==7 and idx==1:
                figure(doc,'role_intervals.png',15.6,'图2  先发与牛棚的过程指标  牛棚估计区间较宽')
            if page==8 and idx==1:
                table(doc,ACTION_TABLE,[2.2,6.4,8.0]);caption(doc,'表5  依据冻结分析提出的阶段安排')
            if page==9 and idx==0:
                table(doc,REPRO_TABLE,[5.7,10.9]);caption(doc,'表6  主要归档内容与复核用途')
    path=OUT/'实习报告_杨炎新_15页.docx';doc.save(path);return path


def executive():
    text=(WORK/'executive_review.md').read_text(encoding='utf-8').split('---',2)[2].strip()
    doc=Document();configure(doc,9,True)
    doc.core_properties.title='今井达也MLB调整建议'
    p=doc.add_paragraph('今井达也MLB调整建议',style='Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0)
    run_font(p.add_run('杨炎新  3230102355     数据窗口  2026年3月29日至8月12日'),8.5,font='微软雅黑')
    lines=text.splitlines();i=0;next_page=False
    while i<len(lines):
        line=lines[i].strip()
        if not line:i+=1;continue
        if line=='[[PAGEBREAK]]':next_page=True;i+=1;continue
        if line.startswith('## '):
            p=doc.add_paragraph(re.sub(r'^[一二三四五六七八九十]+[、.]','',line[3:]),style='Heading 1')
            if next_page:p.paragraph_format.page_break_before=True;next_page=False
        elif line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                row=[c.strip() for c in lines[i].strip().strip('|').split('|')]
                if not all(re.match(r'^:?-+:?$',v) for v in row):rows.append(row)
                i+=1
            widths=[2.1,5.4,9.1] if rows[0][0]=='阶段' else [2.7,6.3,7.6]
            table(doc,rows,widths,size=9,executive=True);continue
        else:ptext(doc,line,9,True)
        i+=1
    path=OUT/'蒋总汇总_今井达也MLB调整建议_2页.docx';doc.save(path);return path


if __name__=='__main__':
    charts()
    for fn in [report,executive]:
        path=fn(); print(path)
